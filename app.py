import streamlit as st

import matplotlib.pyplot as plt
import pandas as pd

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

import analysis


# Page setup

st.set_page_config(
    page_title="Force-Time Analysis",
    page_icon="📈",
    layout="wide"
)

st.title("Force-Time Analysis")

st.write(
    "Upload a raw ForceDecks CSV to perform "
    "jump force-time analysis."
)


# Jump type

jump_type = st.selectbox(
    "Select Jump Type",
    [
        "Countermovement Jump (CMJ)",
        "Squat Jump (SJ)"
    ]
)


# File upload

uploaded_file = st.file_uploader(
    "Upload ForceDecks Raw Data CSV",
    type=["csv"]
)


# Analysis

if uploaded_file is not None:

    try:

        # Select analysis pipeline

        if jump_type == "Countermovement Jump (CMJ)":

            results, processed_df = analysis.analyze_cmj(
                uploaded_file
            )

            st.success("CMJ analysis completed.")

        else:

            results, processed_df = analysis.analyze_sj(
                uploaded_file
            )

            st.success("SJ analysis completed.")

        # Athlete / trial information

        st.subheader("Trial Information")

        col1, col2, col3 = st.columns(3)

        col1.metric(
            "Body Mass",
            f"{results['mass_kg']:.2f} kg"
        )

        col2.metric(
            "Sampling Frequency",
            f"{results['sampling_frequency_hz']:.0f} Hz"
        )

        col3.metric(
            "Contraction Time",
            f"{results['contraction_time_s']:.3f} s"
        )

        # Primary performance metrics

        st.subheader("Performance")

        col1, col2, col3, col4 = st.columns(4)

        col1.metric(
            "Jump Height",
            f"{results['jump_height_m'] * 100:.1f} cm"
        )

        col2.metric(
            "Takeoff Velocity",
            f"{results['takeoff_velocity_m_s']:.3f} m/s"
        )

        col3.metric(
            "Peak Force",
            f"{results['peak_force_n']:.0f} N"
        )

        col4.metric(
            "Peak Power",
            f"{results['peak_propulsive_power_w']:.0f} W"
        )

        # Phase metrics

        if jump_type == "Countermovement Jump (CMJ)":

            st.subheader("CMJ Phases")

            col1, col2, col3 = st.columns(3)

            col1.metric(
                "Unweighting Duration",
                f"{results['unweighting_duration_s']:.3f} s"
            )

            col2.metric(
                "Braking Duration",
                f"{results['braking_duration_s']:.3f} s"
            )

            col3.metric(
                "Propulsive Duration",
                f"{results['propulsive_duration_s']:.3f} s"
            )

        else:

            st.subheader("SJ Phase")

            st.metric(
                "Propulsive Duration",
                f"{results['propulsive_duration_s']:.3f} s"
            )

        # Impulse

        st.subheader("Impulse")

        if jump_type == "Countermovement Jump (CMJ)":

            col1, col2 = st.columns(2)

            col1.metric(
                "Braking Impulse",
                f"{results['braking_impulse_ns']:.2f} N·s"
            )

            col2.metric(
                "Propulsive Impulse",
                f"{results['propulsive_impulse_ns']:.2f} N·s"
            )

        else:

            st.metric(
                "Propulsive Impulse",
                f"{results['propulsive_impulse_ns']:.2f} N·s"
            )

        # RFD

        st.subheader("Rate of Force Development")

        if jump_type == "Countermovement Jump (CMJ)":

            col1, col2 = st.columns(2)

            col1.metric(
                "Peak Braking RFD (20 ms)",
                f"{results['peak_braking_rfd_n_s']:.0f} N/s"
            )

            col2.metric(
                "Peak Propulsive RFD (20 ms)",
                f"{results['peak_propulsive_rfd_n_s']:.0f} N/s"
            )

        else:

            st.metric(
                "Peak Propulsive RFD (20 ms)",
                f"{results['peak_propulsive_rfd_n_s']:.0f} N/s"
            )

        # Landing

        st.subheader("Landing")

        col1, col2 = st.columns(2)

        col1.metric(
            "Peak Landing Force",
            f"{results['peak_landing_force_n']:.0f} N"
        )

        col2.metric(
            "Peak Loading Rate (20 ms)",
            f"{results['peak_landing_rfd_n_s']:.0f} N/s"
        )

        # Force-time curve

        st.subheader("Force-Time Curve")

        onset_idx = results["movement_onset_idx"]
        takeoff_idx = results["takeoff_idx"]
        landing_idx = results["landing_idx"]

        # Display 1 second beyond takeoff so the
        # flight and landing are visible on the plot
        post_takeoff_samples = int(
            round(results["sampling_frequency_hz"] * 1.0)
        )

        plot_end_idx = min(
            takeoff_idx + post_takeoff_samples,
            processed_df.index[-1]
        )

        body_weight = results["body_weight_n"]

        fig, ax = plt.subplots(figsize=(12, 6))

        # CMJ force-time curve

        if jump_type == "Countermovement Jump (CMJ)":

            min_velocity_idx = results["min_velocity_idx"]
            zero_velocity_idx = results["zero_velocity_idx"]

            # Unweighting phase
            unweighting = processed_df.loc[
                onset_idx:min_velocity_idx
            ]

            ax.plot(
                unweighting["Time_Rel"],
                unweighting["Fz_Filtered"],
                color="tab:blue",
                linewidth=2.5,
                label="Unweighting"
            )

            # Braking phase
            braking = processed_df.loc[
                min_velocity_idx:zero_velocity_idx
            ]

            ax.plot(
                braking["Time_Rel"],
                braking["Fz_Filtered"],
                color="tab:orange",
                linewidth=2.5,
                label="Braking"
            )

            # Propulsive phase
            propulsive = processed_df.loc[
                zero_velocity_idx:takeoff_idx
            ]

            ax.plot(
                propulsive["Time_Rel"],
                propulsive["Fz_Filtered"],
                color="tab:green",
                linewidth=2.5,
                label="Propulsion"
            )

            # CMJ event markers
            ax.axvline(
                results["movement_onset_s"],
                color="gray",
                linestyle=":",
                linewidth=1
            )

            ax.axvline(
                results["min_velocity_time_s"],
                color="gray",
                linestyle=":",
                linewidth=1
            )

            ax.axvline(
                results["zero_velocity_time_s"],
                color="gray",
                linestyle=":",
                linewidth=1
            )

            ax.axvline(
                results["takeoff_time_s"],
                color="gray",
                linestyle=":",
                linewidth=1
            )

            ax.set_title(
                "Countermovement Jump Force-Time Curve"
            )

        # SJ force-time curve

        else:

            propulsive = processed_df.loc[
                onset_idx:takeoff_idx
            ]

            ax.plot(
                propulsive["Time_Rel"],
                propulsive["Fz_Filtered"],
                color="tab:green",
                linewidth=2.5,
                label="Propulsion"
            )

            # SJ onset
            ax.axvline(
                results["movement_onset_s"],
                color="gray",
                linestyle=":",
                linewidth=1
            )

            # SJ takeoff
            ax.axvline(
                results["takeoff_time_s"],
                color="gray",
                linestyle=":",
                linewidth=1
            )

            ax.set_title(
                "Squat Jump Force-Time Curve"
            )

        # Flight

        flight = processed_df.loc[
            takeoff_idx:landing_idx
        ]

        ax.plot(
            flight["Time_Rel"],
            flight["Fz_Filtered"],
            color="gray",
            linewidth=2.0,
            label="Flight"
        )

        # Landing

        landing = processed_df.loc[
            landing_idx:plot_end_idx
        ]

        ax.plot(
            landing["Time_Rel"],
            landing["Fz_Filtered"],
            color="tab:red",
            linewidth=2.5,
            label="Landing"
        )

        # Landing contact marker
        ax.axvline(
            results["landing_time_s"],
            color="tab:red",
            linestyle=":",
            linewidth=1.5,
            label="Landing Contact"
        )

        # Body-weight reference

        ax.axhline(
            body_weight,
            color="gray",
            linestyle="--",
            linewidth=1.5,
            label="Body Weight"
        )

        ax.set_xlabel("Time (s)")
        ax.set_ylabel("Vertical Force (N)")

        ax.legend(
            frameon=False,
            ncol=4
        )

        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)

        fig.tight_layout()

        st.pyplot(fig)

        # Raw vs filtered force

        st.subheader("Raw vs Filtered Force")

        raw_filtered = processed_df.loc[
            onset_idx:plot_end_idx
        ]

        fig_raw, ax_raw = plt.subplots(
            figsize=(12, 6)
        )

        # Raw force
        ax_raw.plot(
            raw_filtered["Time_Rel"],
            raw_filtered["Fz"],
            linewidth=1.25,
            alpha=0.55,
            label="Raw Force"
        )

        # 15-Hz filtered force
        ax_raw.plot(
            raw_filtered["Time_Rel"],
            raw_filtered["Fz_Filtered"],
            linewidth=2.5,
            label="15 Hz Filtered Force"
        )

        # 85-Hz landing-impact force
        ax_raw.plot(
            raw_filtered["Time_Rel"],
            raw_filtered["Fz_Landing_Filtered"],
            linewidth=2.0,
            label="85 Hz Filtered Force"
        )

        # Takeoff
        ax_raw.axvline(
            results["takeoff_time_s"],
            color="gray",
            linestyle=":",
            linewidth=1.5,
            label="Takeoff"
        )

        # Landing contact
        ax_raw.axvline(
            results["landing_time_s"],
            color="tab:red",
            linestyle=":",
            linewidth=1.5,
            label="Landing Contact"
        )

        # Body-weight reference
        ax_raw.axhline(
            body_weight,
            color="gray",
            linestyle="--",
            linewidth=1.5,
            label="Body Weight"
        )

        ax_raw.set_title(
            "Raw vs 15 Hz vs 85 Hz Filtered Force"
        )

        ax_raw.set_xlabel("Time (s)")
        ax_raw.set_ylabel("Vertical Force (N)")

        ax_raw.legend(
            frameon=False,
            ncol=3
        )

        ax_raw.spines["top"].set_visible(False)
        ax_raw.spines["right"].set_visible(False)

        fig_raw.tight_layout()

        st.pyplot(fig_raw)

        # Export results

        st.subheader("Export")

        metadata = results["metadata"]

        recording_date = metadata.get(
            "Recording Date", ""
        )

        athlete_id = metadata.get(
            "AthleteId", ""
        )

        # CSV export

        export_results = {
            key: value
            for key, value in results.items()
            if key != "metadata"
        }

        export_results["recording_date"] = (
            recording_date
        )

        export_results["athlete_id"] = (
            athlete_id
        )

        export_results["jump_type"] = (
            "CMJ"
            if jump_type
            == "Countermovement Jump (CMJ)"
            else "SJ"
        )

        export_df = pd.DataFrame(
            [export_results]
        )

        csv = export_df.to_csv(
            index=False
        ).encode("utf-8")

        # Formatted Excel report

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = "Force-Time Report"

        # Excel styles

        title_font = Font(
            size=18,
            bold=True
        )

        section_font = Font(
            size=12,
            bold=True
        )

        header_font = Font(
            bold=True
        )

        thin_border = Border(
            bottom=Side(
                style="thin"
            )
        )

        # Report title

        worksheet["A1"] = (
            "Force-Time Analysis Report"
        )

        worksheet["A1"].font = title_font

        worksheet.merge_cells(
            "A1:C1"
        )

        worksheet["A2"] = (
            "Countermovement Jump"
            if jump_type
            == "Countermovement Jump (CMJ)"
            else "Squat Jump"
        )

        worksheet.merge_cells(
            "A2:C2"
        )

        worksheet["A2"].font = Font(
            italic=True
        )

        # Mutable row tracker used by helper functions
        row_tracker = [4]

        # Excel helper functions

        def add_section(title):

            current_row = row_tracker[0]

            worksheet.cell(
                row=current_row,
                column=1,
                value=title
            )

            worksheet.cell(
                row=current_row,
                column=1
            ).font = section_font

            worksheet.cell(
                row=current_row,
                column=1
            ).border = thin_border

            worksheet.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row,
                end_column=3
            )

            current_row += 1

            worksheet.cell(
                row=current_row,
                column=1,
                value="Metric"
            )

            worksheet.cell(
                row=current_row,
                column=2,
                value="Value"
            )

            worksheet.cell(
                row=current_row,
                column=3,
                value="Unit"
            )

            for column in range(1, 4):

                worksheet.cell(
                    row=current_row,
                    column=column
                ).font = header_font

            current_row += 1

            row_tracker[0] = current_row

        def add_metric(
            metric,
            value,
            unit="",
            decimals=2
        ):

            current_row = row_tracker[0]

            worksheet.cell(
                row=current_row,
                column=1,
                value=metric
            )

            if isinstance(
                value,
                (int, float)
            ):

                cell = worksheet.cell(
                    row=current_row,
                    column=2,
                    value=float(value)
                )

                if decimals == 0:

                    cell.number_format = "0"

                elif decimals == 1:

                    cell.number_format = "0.0"

                elif decimals == 2:

                    cell.number_format = "0.00"

                elif decimals == 3:

                    cell.number_format = "0.000"

            else:

                worksheet.cell(
                    row=current_row,
                    column=2,
                    value=value
                )

            worksheet.cell(
                row=current_row,
                column=3,
                value=unit
            )

            row_tracker[0] = current_row + 1

        def add_space():

            row_tracker[0] += 1

        # Trial information

        add_section(
            "Trial Information"
        )

        add_metric(
            "Jump Type",
            (
                "Countermovement Jump"
                if jump_type
                == "Countermovement Jump (CMJ)"
                else "Squat Jump"
            )
        )

        add_metric(
            "Recording Date",
            recording_date
        )

        add_metric(
            "Athlete ID",
            athlete_id
        )

        add_metric(
            "Body Mass",
            results["mass_kg"],
            "kg",
            2
        )

        add_metric(
            "Sampling Frequency",
            results[
                "sampling_frequency_hz"
            ],
            "Hz",
            0
        )

        add_metric(
            "Contraction Time",
            results[
                "contraction_time_s"
            ],
            "s",
            3
        )

        add_space()

        # Performance

        add_section(
            "Performance"
        )

        add_metric(
            "Jump Height",
            results[
                "jump_height_m"
            ] * 100,
            "cm",
            1
        )

        add_metric(
            "Takeoff Velocity",
            results[
                "takeoff_velocity_m_s"
            ],
            "m/s",
            3
        )

        add_metric(
            "Peak Force",
            results[
                "peak_force_n"
            ],
            "N",
            0
        )

        add_metric(
            "Peak Power",
            results[
                "peak_propulsive_power_w"
            ],
            "W",
            0
        )

        add_space()

        # Phase durations

        add_section(
            "Phase Durations"
        )

        if jump_type == (
            "Countermovement Jump (CMJ)"
        ):

            add_metric(
                "Unweighting Duration",
                results[
                    "unweighting_duration_s"
                ],
                "s",
                3
            )

            add_metric(
                "Braking Duration",
                results[
                    "braking_duration_s"
                ],
                "s",
                3
            )

        add_metric(
            "Propulsive Duration",
            results[
                "propulsive_duration_s"
            ],
            "s",
            3
        )

        add_space()

        # Impulse

        add_section(
            "Impulse"
        )

        if jump_type == (
            "Countermovement Jump (CMJ)"
        ):

            add_metric(
                "Braking Impulse",
                results[
                    "braking_impulse_ns"
                ],
                "N·s",
                2
            )

        add_metric(
            "Propulsive Impulse",
            results[
                "propulsive_impulse_ns"
            ],
            "N·s",
            2
        )

        add_space()

        # RFD

        add_section(
            "Rate of Force Development"
        )

        if jump_type == (
            "Countermovement Jump (CMJ)"
        ):

            add_metric(
                "Peak Braking RFD (20 ms)",
                results[
                    "peak_braking_rfd_n_s"
                ],
                "N/s",
                0
            )

        add_metric(
            "Peak Propulsive RFD (20 ms)",
            results[
                "peak_propulsive_rfd_n_s"
            ],
            "N/s",
            0
        )

        add_space()

        # Landing

        add_section(
            "Landing"
        )

        add_metric(
            "Peak Landing Force",
            results[
                "peak_landing_force_n"
            ],
            "N",
            0
        )

        add_metric(
            "Peak Loading Rate (20 ms)",
            results[
                "peak_landing_rfd_n_s"
            ],
            "N/s",
            0
        )

        add_space()

        # Processing information

        add_section(
            "Processing Information"
        )

        add_metric(
            "Movement Force Filter",
            results.get(
                "filter_cutoff_hz",
                15.0
            ),
            "Hz",
            0
        )

        add_metric(
            "Landing Impact Filter",
            results.get(
                "landing_filter_cutoff_hz",
                85.0
            ),
            "Hz",
            0
        )

        add_metric(
            "Filter Order",
            results.get(
                "filter_order",
                4
            ),
            "",
            0
        )

        add_metric(
            "RFD / Loading Rate Window",
            results.get(
                "rfd_window_ms",
                20.0
            ),
            "ms",
            0
        )

        add_metric(
            "Body Weight",
            results[
                "body_weight_n"
            ],
            "N",
            2
        )

        if jump_type == (
            "Countermovement Jump (CMJ)"
        ):

            add_metric(
                "Takeoff Threshold",
                results.get(
                    "takeoff_threshold_n",
                    ""
                ),
                "N",
                2
            )

        add_metric(
            "Landing Threshold",
            results.get(
                "landing_threshold_n",
                ""
            ),
            "N",
            2
        )

        # Final Excel formatting

        worksheet.column_dimensions[
            "A"
        ].width = 34

        worksheet.column_dimensions[
            "B"
        ].width = 20

        worksheet.column_dimensions[
            "C"
        ].width = 14

        for row in worksheet.iter_rows():

            for cell in row:

                cell.alignment = Alignment(
                    vertical="center"
                )

        worksheet.freeze_panes = "A4"

        # Create Excel file in memory

        excel_buffer = BytesIO()

        workbook.save(
            excel_buffer
        )

        excel_buffer.seek(0)

        # Export file names

        if jump_type == (
            "Countermovement Jump (CMJ)"
        ):

            csv_filename = (
                "CMJ_analysis_results.csv"
            )

            excel_filename = (
                "CMJ_force_time_report.xlsx"
            )

        else:

            csv_filename = (
                "SJ_analysis_results.csv"
            )

            excel_filename = (
                "SJ_force_time_report.xlsx"
            )

        # Download buttons

        col1, col2 = st.columns(2)

        with col1:

            st.download_button(
                label="Download Results CSV",
                data=csv,
                file_name=csv_filename,
                mime="text/csv",
                use_container_width=True
            )

        with col2:

            st.download_button(
                label="Download Formatted Excel Report",
                data=excel_buffer.getvalue(),
                file_name=excel_filename,
                mime=(
                    "application/"
                    "vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                use_container_width=True
            )

    except Exception as error:

        st.error(
            f"Analysis failed: {error}"
        )